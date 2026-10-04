import argparse, contextlib, csv, gc, re
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pytorch_lightning as pl
import torch
import torch.nn.functional as F
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from torch import nn
from torch.utils.data import DataLoader, Dataset
from tqdm import tqdm

from meteolip import create_model_and_transforms
from meteolip_train.data import get_classification_dataset


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--model", required=True)
    p.add_argument("--checkpoint", required=True)
    p.add_argument("--train-classification-json-path", required=True)
    p.add_argument("--test-classification-json-path", required=True)
    p.add_argument("--classification-npz-dir", required=True)
    p.add_argument("--save-dir", required=True)
    p.add_argument("--batch-size", type=int, default=64)
    p.add_argument("--probe-batch-size", type=int, default=256)
    p.add_argument("--epochs", type=int, default=20)
    p.add_argument("--lr", type=float, default=1e-3)
    p.add_argument("--weight-decay", type=float, default=0.01)
    p.add_argument("--workers", type=int, default=4)
    p.add_argument("--seeds", type=int, nargs="+", default=[0])
    p.add_argument("--device", choices=("cpu", "cuda"), default="cuda")
    p.add_argument("--precision", choices=("fp32", "amp"), default="fp32")
    return p.parse_args()


def load_cpu(path):
    try:
        return torch.load(path, map_location="cpu", weights_only=False)
    except TypeError:
        return torch.load(path, map_location="cpu")


def frozen_vision(model_name, checkpoint_path):
    model, _, _ = create_model_and_transforms(
        model_name, pretrained=None, device="cpu", precision="fp32"
    )
    ckpt = load_cpu(checkpoint_path)
    state = ckpt.get("state_dict", ckpt.get("model", ckpt)) if isinstance(ckpt, dict) else ckpt
    if not isinstance(state, dict):
        raise TypeError("Checkpoint does not contain a state dict.")

    visual_state = {}
    for raw_key, value in state.items():
        key = raw_key.removeprefix("module.").removeprefix("model.")
        if not key.startswith("visual."):
            continue
        key = key[len("visual."):]
        if key in visual_state:
            raise RuntimeError(f"Duplicate visual key after prefix cleanup: {key}")
        visual_state[key] = value
    
    visual_state = {key.replace("vat", "mettok"): val for key, val in visual_state.items()}

    if not visual_state:
        raise RuntimeError("No visual.* tensors found in checkpoint.")
    model.visual.load_state_dict(visual_state, strict=True)
    print(f"Strictly loaded {len(visual_state)} visual tensors.")
    vision = model.visual.requires_grad_(False).eval()
    del model, ckpt, state, visual_state
    return vision


def as_feature_tensor(output):
    if isinstance(output, dict):
        output = output["image_features"]
    elif isinstance(output, (tuple, list)):
        output = output[0]
    if not isinstance(output, torch.Tensor) or output.ndim != 2:
        raise ValueError("Vision output must be a [batch, feature_dim] tensor.")
    return output


def amp_ctx(device, precision):
    if precision == "fp32":
        return contextlib.nullcontext()
    if device.type != "cuda":
        raise ValueError("--precision amp requires --device cuda.")
    return torch.autocast("cuda", dtype=torch.float16)


def eval_loader(loader):
    return DataLoader(
        loader.dataset,
        batch_size=loader.batch_size,
        shuffle=False,
        num_workers=loader.num_workers,
        collate_fn=loader.collate_fn,
        pin_memory=loader.pin_memory,
        persistent_workers=bool(loader.num_workers and loader.persistent_workers),
        drop_last=False,
    )


@torch.inference_mode()
def extract_image_features(vision, loader, device, precision):
    xs, ys = [], []
    vision = vision.to(device).eval()
    for batch in tqdm(eval_loader(loader)):
        images, target = batch[:2]
        images = images.to(device=device, dtype=torch.float32, non_blocking=True)
        with amp_ctx(device, precision):
            xs.append(as_feature_tensor(vision(images)).detach().float().cpu())
        ys.append(target.detach().long().flatten().cpu())
    if not xs:
        raise ValueError("Empty dataloader.")
    return torch.cat(xs), torch.cat(ys)


def l2_normalize(x, name, eps=1e-12):
    if x.ndim != 2:
        raise ValueError(f"{name}: expected [N, D], got {tuple(x.shape)}")
    if (x.norm(p=2, dim=1) <= eps).any():
        raise ValueError(f"{name}: found zero-length feature vector before normalization.")
    x = F.normalize(x, p=2, dim=1, eps=eps)
    norms = x.norm(p=2, dim=1)
    if not torch.allclose(norms, torch.ones_like(norms), rtol=1e-5, atol=1e-5):
        raise RuntimeError(f"{name}: F.normalize(..., dim=1) did not produce unit norms.")
    return x


def check_splits(train_loaders, test_loaders):
    if not train_loaders or set(train_loaders) != set(test_loaders):
        raise ValueError(f"Task mismatch: train={sorted(train_loaders)}, test={sorted(test_loaders)}")
    task_classes = {}
    for task in train_loaders:
        train_names = list(train_loaders[task].dataset.classes)
        test_names = list(test_loaders[task].dataset.classes)
        if train_names != test_names:
            raise ValueError(f"Class/order mismatch for {task}:\ntrain={train_names}\ntest ={test_names}")
        if not train_names or len(train_names) != len(set(train_names)):
            raise ValueError(f"Task {task} must have non-empty, unique class names.")
        task_classes[task] = train_names
    return task_classes


def validate_xy(x, y, num_classes, name):
    if x.ndim != 2 or y.ndim != 1 or len(x) != len(y):
        raise ValueError(f"{name}: invalid feature/target shapes: {tuple(x.shape)}, {tuple(y.shape)}")
    if not torch.isfinite(x).all():
        raise ValueError(f"{name}: non-finite features found.")
    if y.min().item() < 0 or y.max().item() >= num_classes:
        raise ValueError(f"{name}: labels are outside [0, {num_classes}).")


class FeatureDataset(Dataset):
    def __init__(self, x, y):
        validate_xy(x, y, int(y.max().item()) + 1, "FeatureDataset")
        self.x, self.y = x.float().contiguous(), y.long().contiguous()

    def __len__(self):
        return len(self.y)

    def __getitem__(self, i):
        return self.x[i], self.y[i]


def feature_loader(dataset, batch_size, shuffle, seed, pin_memory):
    g = torch.Generator().manual_seed(seed)
    return DataLoader(
        dataset,
        batch_size=batch_size,
        shuffle=shuffle,
        generator=g if shuffle else None,
        num_workers=0,
        pin_memory=pin_memory,
    )


class LinearProbe(pl.LightningModule):
    def __init__(self, dim, num_classes, lr, weight_decay, task, seed):
        super().__init__()
        self.head = nn.Linear(dim, num_classes)
        self.num_classes, self.lr, self.weight_decay = num_classes, lr, weight_decay
        self.task, self.seed, self.history = task, seed, []
        self.last_confusion_matrix = None

    def forward(self, x):
        return self.head(x)

    def on_train_epoch_start(self):
        self.train_loss_sum = torch.zeros((), device=self.device)
        self.train_count = 0

    def training_step(self, batch, batch_idx):
        x, y = batch
        loss = F.cross_entropy(self(x), y)
        self.train_loss_sum += loss.detach() * y.numel()
        self.train_count += y.numel()
        return loss

    def on_validation_epoch_start(self):
        self.test_loss_sum = torch.zeros((), device=self.device)
        self.test_count = 0
        self.correct1 = torch.zeros((), dtype=torch.long, device=self.device)
        self.correct3 = torch.zeros((), dtype=torch.long, device=self.device)
        self.cm = torch.zeros(self.num_classes, self.num_classes, dtype=torch.long, device=self.device)

    def validation_step(self, batch, batch_idx):
        x, y = batch
        logits = self(x)
        pred = logits.argmax(1)
        top3 = logits.topk(min(3, self.num_classes), dim=1).indices
        loss = F.cross_entropy(logits, y)
        self.test_loss_sum += loss.detach() * y.numel()
        self.test_count += y.numel()
        self.correct1 += pred.eq(y).sum()
        self.correct3 += top3.eq(y[:, None]).any(1).sum()
        idx = y * self.num_classes + pred
        self.cm += torch.bincount(idx, minlength=self.num_classes**2).reshape(
            self.num_classes, self.num_classes
        )

    def on_validation_epoch_end(self):
        train_loss = self.train_loss_sum / max(1, self.train_count)
        test_loss = self.test_loss_sum / max(1, self.test_count)
        cm = self.cm.float()
        tp, denom = cm.diag(), cm.sum(0) + cm.sum(1)
        f1 = torch.where(denom > 0, 2 * tp / denom, 0.0).mean()
        top1 = self.correct1.float() / max(1, self.test_count)
        top3 = self.correct3.float() / max(1, self.test_count)
        self.last_confusion_matrix = self.cm.detach().cpu().clone()
        row = dict(
            epoch=int(self.current_epoch + 1),
            train_loss=float(train_loss.cpu()),
            test_loss=float(test_loss.cpu()),
            test_macro_f1=float(f1.cpu()),
            test_acc_top1=float(top1.cpu()),
            test_acc_top3=float(top3.cpu()),
        )
        self.history.append(row)
        self.log_dict(
            {
                "train_loss": train_loss,
                "test_loss": test_loss,
                "test_macro_f1": f1,
                "test_acc_top1": top1,
                "test_acc_top3": top3,
            },
            prog_bar=True,
        )
        print(
            f"{self.task} | seed={self.seed} | epoch={row['epoch']:03d} "
            f"train_loss={row['train_loss']:.6f} test_loss={row['test_loss']:.6f} "
            f"f1={row['test_macro_f1']:.4f} top1={row['test_acc_top1']:.4f} "
            f"top3={row['test_acc_top3']:.4f}"
        )

    def configure_optimizers(self):
        return torch.optim.AdamW(self.head.parameters(), lr=self.lr, weight_decay=self.weight_decay)


def trainer_args(device, precision):
    if device == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("--device cuda requested, but CUDA is unavailable.")
    if precision == "amp" and device != "cuda":
        raise ValueError("--precision amp requires --device cuda.")
    return dict(
        accelerator="gpu" if device == "cuda" else "cpu",
        devices=1,
        precision="16-mixed" if precision == "amp" else "32-true",
        logger=False,
        enable_checkpointing=False,
        enable_model_summary=False,
        num_sanity_val_steps=0,
        check_val_every_n_epoch=1,
    )


def safe_name(name):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("._") or "task"


def safe_sheet(name):
    return re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "sheet"


def style(ws):
    fill, font = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)
    for c in ws[1]:
        c.fill, c.font = fill, font
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = min(
            max(max(len(str(c.value or "")) for c in col) + 2, 12), 32
        )
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = "0.0000"


def write_excel(path, summary, histories):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Task", "Seed", "Macro F1", "Acc Top-1", "Acc Top-3"])
    for r in summary:
        ws.append([r["task"], r["seed"], r["macro_f1"], r["acc_top1"], r["acc_top3"]])
    style(ws)

    for seed, rows in histories.items():
        ws = wb.create_sheet(safe_sheet(f"seed_{seed}"))
        ws.append(
            [
                "Task",
                "Epoch",
                "Train Loss",
                "Test Loss",
                "Test Macro F1",
                "Test Acc Top-1",
                "Test Acc Top-3",
            ]
        )
        for r in rows:
            ws.append(
                [
                    r["task"],
                    r["epoch"],
                    r["train_loss"],
                    r["test_loss"],
                    r["test_macro_f1"],
                    r["test_acc_top1"],
                    r["test_acc_top3"],
                ]
            )
        style(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"Saved Excel results to: {path}")


def save_linear_head(path, probe, class_names, feature_dim, final_metrics):
    path.parent.mkdir(parents=True, exist_ok=True)
    torch.save(
        {
            "task": probe.task,
            "seed": probe.seed,
            "epoch": final_metrics["epoch"],
            "feature_dim": feature_dim,
            "num_classes": len(class_names),
            "class_names": list(class_names),
            "feature_normalization": "F.normalize(features, p=2, dim=1)",
            "metrics": dict(final_metrics),
            "linear_head_state_dict": {
                key: value.detach().cpu().clone()
                for key, value in probe.head.state_dict().items()
            },
        },
        path,
    )
    print(f"Saved last-epoch linear head to: {path}")


def save_confusion_matrix(cm, class_names, png_path, title):
    if cm is None or tuple(cm.shape) != (len(class_names), len(class_names)):
        raise ValueError("Invalid last-epoch confusion matrix.")

    png_path.parent.mkdir(parents=True, exist_ok=True)
    cm_np = cm.numpy()
    size = min(24, max(7, 0.65 * len(class_names)))
    fig, ax = plt.subplots(figsize=(size, size))
    image = ax.imshow(cm_np, interpolation="nearest", cmap="Blues")
    fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)
    ax.set_title(title)
    ax.set_xlabel("Predicted class")
    ax.set_ylabel("True class")
    ax.set_xticks(range(len(class_names)), labels=class_names, rotation=45, ha="right")
    ax.set_yticks(range(len(class_names)), labels=class_names)

    threshold = float(cm_np.max()) / 2.0 if cm_np.size else 0.0
    font_size = max(5, min(10, 180 // max(1, len(class_names))))
    for true_idx in range(len(class_names)):
        for pred_idx in range(len(class_names)):
            value = int(cm_np[true_idx, pred_idx])
            ax.text(
                pred_idx,
                true_idx,
                str(value),
                ha="center",
                va="center",
                fontsize=font_size,
                color="white" if value > threshold else "black",
            )

    fig.tight_layout()
    fig.savefig(png_path, dpi=300, bbox_inches="tight")
    plt.close(fig)

    csv_path = png_path.with_suffix(".csv")
    with csv_path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(["true\\pred", *class_names])
        for class_name, row in zip(class_names, cm.tolist()):
            writer.writerow([class_name, *row])

    print(f"Saved last-epoch confusion matrix to: {png_path}")
    print(f"Saved raw confusion-matrix counts to: {csv_path}")


def main():
    args = parse_args()
    if len(set(args.seeds)) != len(args.seeds):
        raise ValueError(f"Duplicate seeds are not allowed: {args.seeds}")
    if args.precision == "fp32":
        torch.set_float32_matmul_precision("highest")
        if torch.cuda.is_available():
            torch.backends.cuda.matmul.allow_tf32 = False
            torch.backends.cudnn.allow_tf32 = False

    save_dir = Path(args.save_dir)
    save_dir.mkdir(parents=True, exist_ok=True)
    device = torch.device(args.device)
    train_loaders = get_classification_dataset(
        args.train_classification_json_path,
        args.classification_npz_dir,
        args.batch_size,
        args.workers,
    )
    test_loaders = get_classification_dataset(
        args.test_classification_json_path,
        args.classification_npz_dir,
        args.batch_size,
        args.workers,
    )
    task_classes = check_splits(train_loaders, test_loaders)
    vision = frozen_vision(args.model, args.checkpoint)
    common = trainer_args(args.device, args.precision)
    summary, histories, visual_dim = [], defaultdict(list), None

    for task, class_names in task_classes.items():
        print(f"\nTask: {task} | {len(class_names)} classes={class_names}")
        train_x, train_y = extract_image_features(
            vision, train_loaders[task], device, args.precision
        )
        test_x, test_y = extract_image_features(
            vision, test_loaders[task], device, args.precision
        )
        validate_xy(train_x, train_y, len(class_names), f"{task}/train")
        validate_xy(test_x, test_y, len(class_names), f"{task}/test")
        if train_x.shape[1] != test_x.shape[1]:
            raise ValueError(f"{task}: train/test embedding dimensions differ.")
        visual_dim = visual_dim or int(train_x.shape[1])
        if train_x.shape[1] != visual_dim:
            raise ValueError(f"{task}: embedding dim differs from previous tasks.")
        print(f"Visual embedding dim: {visual_dim}")

        train_x = l2_normalize(train_x, f"{task}/train")
        test_x = l2_normalize(test_x, f"{task}/test")
        print(
            "Applied F.normalize(features, p=2, dim=1); "
            "each image feature row now has L2 norm 1."
        )

        train_ds, test_ds = FeatureDataset(train_x, train_y), FeatureDataset(test_x, test_y)
        task_dir = safe_name(task)
        for seed in args.seeds:
            pl.seed_everything(seed, workers=True)
            train_loader = feature_loader(
                train_ds, args.probe_batch_size, True, seed, args.device == "cuda"
            )
            test_loader = feature_loader(
                test_ds, args.probe_batch_size, False, seed, args.device == "cuda"
            )
            probe = LinearProbe(
                visual_dim, len(class_names), args.lr, args.weight_decay, task, seed
            )
            trainer = pl.Trainer(**common, max_epochs=args.epochs)
            trainer.fit(
                probe,
                train_dataloaders=train_loader,
                val_dataloaders=test_loader,
            )
            if len(probe.history) != args.epochs:
                raise RuntimeError(
                    f"Expected {args.epochs} epoch records, got {len(probe.history)}"
                )

            final = probe.history[-1]
            summary.append(
                dict(
                    task=task,
                    seed=seed,
                    macro_f1=final["test_macro_f1"],
                    acc_top1=final["test_acc_top1"],
                    acc_top3=final["test_acc_top3"],
                )
            )
            histories[seed].extend(dict(task=task, **row) for row in probe.history)

            run_name = f"seed_{seed}_epoch_{final['epoch']}"
            head_path = save_dir / "linear_heads" / task_dir / f"{run_name}.pt"
            cm_path = save_dir / "confusion_matrices" / task_dir / f"{run_name}.png"
            save_linear_head(head_path, probe, class_names, visual_dim, final)
            save_confusion_matrix(
                probe.last_confusion_matrix,
                class_names,
                cm_path,
                title=f"{task} | seed={seed} | epoch={final['epoch']}",
            )

            print(
                f"FINAL {task} | seed={seed} | last_epoch={final['epoch']} "
                f"macro_f1={final['test_macro_f1']:.4f} "
                f"acc_top1={final['test_acc_top1']:.4f} "
                f"acc_top3={final['test_acc_top3']:.4f}"
            )
            del trainer, probe, train_loader, test_loader
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()

        del train_ds, test_ds, train_x, train_y, test_x, test_y
        gc.collect()

    del vision, train_loaders, test_loaders
    write_excel(save_dir / "linear_probe_results.xlsx", summary, dict(histories))


if __name__ == "__main__":
    main()
