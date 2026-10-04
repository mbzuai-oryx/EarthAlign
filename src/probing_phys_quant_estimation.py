import argparse, contextlib, csv, gc, math, re
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pytorch_lightning as pl
import torch
import torch.nn.functional as F
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from torch import nn
from torch.utils.data import DataLoader, Dataset
from tqdm import tqdm

from meteolip import create_model_and_transforms
from meteolip_train.data import get_ibtracs_reg_dataset
from meteolip_train.zero_shot import compute_regression_metrics


TASK_PATH_ARGS = {
    "wmo_wind": ("train_wind_json_path", "test_wind_json_path"),
    "wmo_pres": ("train_pres_json_path", "test_pres_json_path"),
}


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--model", required=True)
    p.add_argument("--checkpoint", required=True)
    p.add_argument("--train-wind-json-path", required=True)
    p.add_argument("--train-pres-json-path", required=True)
    p.add_argument("--test-wind-json-path", required=True)
    p.add_argument("--test-pres-json-path", required=True)
    p.add_argument("--ibtracs-npz-dir", required=True)
    p.add_argument("--save-dir", required=True)
    p.add_argument("--batch-size", type=int, default=64)
    p.add_argument("--probe-batch-size", type=int, default=256)
    p.add_argument("--epochs", type=int, default=20)
    p.add_argument("--lr", type=float, default=1e-3)
    p.add_argument("--weight-decay", type=float, default=0.01)
    p.add_argument("--workers", type=int, default=4)
    p.add_argument("--seeds", type=int, nargs="+", default=[0])
    p.add_argument("--device", choices=("cpu", "cuda"), default="cuda")
    p.add_argument("--precision", choices=("fp32", "amp"), default="fp32")
    return p.parse_args()


def load_cpu(path):
    try:
        return torch.load(path, map_location="cpu", weights_only=False)
    except TypeError:
        return torch.load(path, map_location="cpu")


def frozen_vision(model_name, checkpoint_path):
    model, _, _ = create_model_and_transforms(
        model_name, pretrained=None, device="cpu", precision="fp32"
    )
    ckpt = load_cpu(checkpoint_path)
    state = ckpt.get("state_dict", ckpt.get("model", ckpt)) if isinstance(ckpt, dict) else ckpt
    if not isinstance(state, dict):
        raise TypeError("Checkpoint does not contain a state dict.")

    visual_state = {}
    for raw_key, value in state.items():
        key = raw_key.removeprefix("module.").removeprefix("model.")
        if not key.startswith("visual."):
            continue
        key = key[len("visual."):]
        if key in visual_state:
            raise RuntimeError(f"Duplicate visual key after prefix cleanup: {key}")
        visual_state[key] = value

    visual_state = {key.replace("vat", "mettok"): val for key, val in visual_state.items()}

    if not visual_state:
        raise RuntimeError("No visual.* tensors found in checkpoint.")
    model.visual.load_state_dict(visual_state, strict=True)
    print(f"Strictly loaded {len(visual_state)} visual tensors.")
    vision = model.visual.requires_grad_(False).eval()
    del model, ckpt, state, visual_state
    return vision


def as_feature_tensor(output):
    if isinstance(output, dict):
        output = output["image_features"]
    elif isinstance(output, (tuple, list)):
        output = output[0]
    if not isinstance(output, torch.Tensor) or output.ndim != 2:
        raise ValueError("Vision output must be a [batch, feature_dim] tensor.")
    return output


def amp_ctx(device, precision):
    if precision == "fp32":
        return contextlib.nullcontext()
    if device.type != "cuda":
        raise ValueError("--precision amp requires --device cuda.")
    return torch.autocast("cuda", dtype=torch.float16)


def unwrap_loader(data):
    loader = getattr(data, "dataloader", data)
    if not isinstance(loader, DataLoader):
        raise TypeError("get_ibtracs_reg_dataset must return a DataLoader or an object with .dataloader.")
    return loader


def eval_loader(loader):
    return DataLoader(
        loader.dataset,
        batch_size=loader.batch_size,
        shuffle=False,
        num_workers=loader.num_workers,
        collate_fn=loader.collate_fn,
        pin_memory=loader.pin_memory,
        persistent_workers=bool(loader.num_workers and loader.persistent_workers),
        drop_last=False,
    )


@torch.inference_mode()
def extract_image_features(vision, loader, device, precision, description):
    xs, ys = [], []
    vision = vision.to(device).eval()

    for batch in tqdm(eval_loader(loader), desc=description):
        if not isinstance(batch, (tuple, list)) or len(batch) < 2:
            raise ValueError("Regression batches must contain images and numerical targets.")

        images, targets = batch[0], batch[-1]
        images = images.to(device=device, dtype=torch.float32, non_blocking=True)
        targets = torch.as_tensor(targets, dtype=torch.float32).flatten().cpu()

        with amp_ctx(device, precision):
            features = as_feature_tensor(vision(images))

        features = features.detach().float().cpu()
        if len(features) != len(targets):
            raise ValueError(
                f"Batch feature/target count mismatch: {len(features)} vs {len(targets)}."
            )
        xs.append(features)
        ys.append(targets)

    if not xs:
        raise ValueError("Empty dataloader.")
    return torch.cat(xs), torch.cat(ys)


def l2_normalize(x, name, eps=1e-12):
    if x.ndim != 2:
        raise ValueError(f"{name}: expected [N, D], got {tuple(x.shape)}")
    if (x.norm(p=2, dim=1) <= eps).any():
        raise ValueError(f"{name}: found zero-length feature vector before normalization.")
    x = F.normalize(x.float(), p=2, dim=1, eps=eps)
    norms = x.norm(p=2, dim=1)
    if not torch.allclose(norms, torch.ones_like(norms), rtol=1e-5, atol=1e-5):
        raise RuntimeError(f"{name}: F.normalize(..., dim=1) did not produce unit norms.")
    return x.contiguous()


def validate_xy(x, y, name):
    if x.ndim != 2 or y.ndim != 1 or len(x) != len(y):
        raise ValueError(f"{name}: invalid feature/target shapes: {tuple(x.shape)}, {tuple(y.shape)}")
    if len(y) < 2:
        raise ValueError(f"{name}: at least two samples are required.")
    if not torch.isfinite(x).all():
        raise ValueError(f"{name}: non-finite features found.")
    if not torch.isfinite(y).all():
        raise ValueError(f"{name}: non-finite targets found.")


def target_stats(train_y, task, eps=1e-12):
    mean = train_y.float().mean()
    std = train_y.float().std(unbiased=False)
    if not torch.isfinite(mean) or not torch.isfinite(std):
        raise ValueError(f"{task}: target mean/std is non-finite.")
    if std <= eps:
        raise ValueError(f"{task}: training targets have zero standard deviation.")
    return float(mean), float(std)


class RegressionFeatureDataset(Dataset):
    def __init__(self, x, y):
        validate_xy(x, y, "RegressionFeatureDataset")
        self.x = x.float().contiguous()
        self.y = y.float().contiguous()

    def __len__(self):
        return len(self.y)

    def __getitem__(self, index):
        return self.x[index], self.y[index]


def feature_loader(dataset, batch_size, shuffle, seed, pin_memory):
    generator = torch.Generator().manual_seed(seed)
    return DataLoader(
        dataset,
        batch_size=batch_size,
        shuffle=shuffle,
        generator=generator if shuffle else None,
        num_workers=0,
        pin_memory=pin_memory,
    )



REGRESSION_METRIC_KEYS = ("mae", "rmse", "r2", "pearson")


class LinearRegressionProbe(pl.LightningModule):
    def __init__(self, dim, target_mean, target_std, lr, weight_decay, task, seed):
        super().__init__()
        self.head = nn.Linear(dim, 1)
        self.lr, self.weight_decay = lr, weight_decay
        self.task, self.seed, self.history = task, seed, []
        self.register_buffer("target_mean", torch.tensor(target_mean, dtype=torch.float32))
        self.register_buffer("target_std", torch.tensor(target_std, dtype=torch.float32))
        self.last_predictions = None
        self.last_targets = None

    def forward(self, x):
        return self.head(x).squeeze(1)

    def standardize(self, y):
        return (y.float() - self.target_mean) / self.target_std

    def unstandardize(self, y_standardized):
        return y_standardized.float() * self.target_std + self.target_mean

    def on_train_epoch_start(self):
        self.train_loss_sum = torch.zeros((), device=self.device)
        self.train_count = 0

    def training_step(self, batch, batch_idx):
        x, target = batch
        prediction_standardized = self(x)
        target_standardized = self.standardize(target)
        loss = F.mse_loss(prediction_standardized, target_standardized)
        self.train_loss_sum += loss.detach() * target.numel()
        self.train_count += target.numel()
        return loss

    def on_validation_epoch_start(self):
        self.test_loss_sum = torch.zeros((), device=self.device)
        self.test_count = 0
        self.validation_predictions = []
        self.validation_targets = []

    def validation_step(self, batch, batch_idx):
        x, target = batch
        prediction_standardized = self(x)
        target_standardized = self.standardize(target)
        loss = F.mse_loss(prediction_standardized, target_standardized)

        prediction = self.unstandardize(prediction_standardized)
        self.test_loss_sum += loss.detach() * target.numel()
        self.test_count += target.numel()
        self.validation_predictions.append(prediction.detach().float().cpu())
        self.validation_targets.append(target.detach().float().cpu())

    def on_validation_epoch_end(self):
        if not self.validation_predictions:
            raise RuntimeError("Validation produced no predictions.")

        predictions = torch.cat(self.validation_predictions)
        targets = torch.cat(self.validation_targets)
        all_metrics = compute_regression_metrics(predictions, targets)
        metrics = {
            key: float(all_metrics[key])
            for key in REGRESSION_METRIC_KEYS
        }
        train_loss = self.train_loss_sum / max(1, self.train_count)
        test_loss = self.test_loss_sum / max(1, self.test_count)

        self.last_predictions = predictions.clone()
        self.last_targets = targets.clone()
        row = {
            "epoch": int(self.current_epoch + 1),
            "train_loss": float(train_loss.cpu()),
            "test_loss": float(test_loss.cpu()),
            "test_mae": metrics["mae"],
            "test_rmse": metrics["rmse"],
            "test_r2": metrics["r2"],
            "test_pearson": metrics["pearson"],
        }
        self.history.append(row)
        self.log_dict(
            {
                "train_loss": train_loss,
                "test_loss": test_loss,
                "test_mae": metrics["mae"],
                "test_rmse": metrics["rmse"],
                "test_r2": metrics["r2"],
                "test_pearson": metrics["pearson"],
            },
            prog_bar=True,
        )
        print(
            f"{self.task} | seed={self.seed} | epoch={row['epoch']:03d} "
            f"train_loss={row['train_loss']:.6f} test_loss={row['test_loss']:.6f} "
            f"mae={row['test_mae']:.4f} rmse={row['test_rmse']:.4f} "
            f"r2={row['test_r2']:.4f} pearson={row['test_pearson']:.4f}"
        )

    def configure_optimizers(self):
        return torch.optim.AdamW(self.head.parameters(), lr=self.lr, weight_decay=self.weight_decay)


def trainer_args(device, precision):
    if device == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("--device cuda requested, but CUDA is unavailable.")
    if precision == "amp" and device != "cuda":
        raise ValueError("--precision amp requires --device cuda.")
    return {
        "accelerator": "gpu" if device == "cuda" else "cpu",
        "devices": 1,
        "precision": "16-mixed" if precision == "amp" else "32-true",
        "logger": False,
        "enable_checkpointing": False,
        "enable_model_summary": False,
        "num_sanity_val_steps": 0,
        "check_val_every_n_epoch": 1,
    }


def safe_name(name):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("._") or "task"


def safe_sheet(name):
    return re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "sheet"


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill, cell.font = fill, font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(
            max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 32
        )
        for cell in column[1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"


def excel_value(value):
    return None if isinstance(value, float) and not math.isfinite(value) else value


def write_excel(path, summary, histories):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Task", "Seed", "MAE", "RMSE", "R2", "Pearson"])
    for row in summary:
        ws.append(
            [
                row["task"],
                row["seed"],
                excel_value(row["mae"]),
                excel_value(row["rmse"]),
                excel_value(row["r2"]),
                excel_value(row["pearson"]),
            ]
        )
    style_sheet(ws)

    for seed, rows in histories.items():
        ws = wb.create_sheet(safe_sheet(f"seed_{seed}"))
        ws.append(
            [
                "Task",
                "Epoch",
                "Train Loss (standardized MSE)",
                "Test Loss (standardized MSE)",
                "Test MAE",
                "Test RMSE",
                "Test R2",
                "Test Pearson",
            ]
        )
        for row in rows:
            ws.append(
                [
                    row["task"],
                    row["epoch"],
                    excel_value(row["train_loss"]),
                    excel_value(row["test_loss"]),
                    excel_value(row["test_mae"]),
                    excel_value(row["test_rmse"]),
                    excel_value(row["test_r2"]),
                    excel_value(row["test_pearson"]),
                ]
            )
        style_sheet(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"Saved Excel results to: {path}")


def save_linear_head(path, probe, feature_dim, final_metrics):
    path.parent.mkdir(parents=True, exist_ok=True)
    torch.save(
        {
            "task": probe.task,
            "seed": probe.seed,
            "epoch": final_metrics["epoch"],
            "feature_dim": feature_dim,
            "output_dim": 1,
            "feature_normalization": "F.normalize(features, p=2, dim=1)",
            "target_standardization": {
                "mean": float(probe.target_mean.cpu()),
                "std": float(probe.target_std.cpu()),
                "std_unbiased": False,
            },
            "losses_standardized_mse": {
                "train": final_metrics["train_loss"],
                "test": final_metrics["test_loss"],
            },
            "metrics_original_units": {
                "mae": final_metrics["test_mae"],
                "rmse": final_metrics["test_rmse"],
                "r2": final_metrics["test_r2"],
                "pearson": final_metrics["test_pearson"],
            },
            "linear_head_state_dict": {
                key: value.detach().cpu().clone()
                for key, value in probe.head.state_dict().items()
            },
        },
        path,
    )
    print(f"Saved last-epoch linear head to: {path}")


def metric_text(metrics):
    return (
        f"MAE: {metrics['mae']:.4f}\n"
        f"RMSE: {metrics['rmse']:.4f}\n"
        f"R²: {metrics['r2']:.4f}\n"
        f"Pearson: {metrics['pearson']:.4f}"
    )


def save_regression_plot(predictions, targets, metrics, png_path, task, title):
    predictions = predictions.flatten().float().cpu()
    targets = targets.flatten().float().cpu()
    if len(predictions) != len(targets) or not len(targets):
        raise ValueError("Invalid predictions/targets for regression plot.")

    low = min(float(predictions.min()), float(targets.min()))
    high = max(float(predictions.max()), float(targets.max()))
    span = high - low
    pad = 0.05 * span if span > 0 else max(1.0, abs(low) * 0.05)
    limits = [low - pad, high + pad]

    png_path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(7, 7))
    ax.scatter(targets.numpy(), predictions.numpy(), alpha=0.5, s=18)
    ax.plot(limits, limits, linestyle="--", linewidth=1.5, label="Perfect prediction")
    ax.set_xlim(limits)
    ax.set_ylim(limits)
    ax.set_aspect("equal", adjustable="box")
    label = {"wmo_wind": "WMO wind", "wmo_pres": "WMO pressure"}.get(task, task)
    ax.set_xlabel(f"Actual {label}")
    ax.set_ylabel(f"Predicted {label}")
    ax.set_title(title)
    ax.grid(True, alpha=0.3)
    ax.legend(loc="lower right")
    ax.text(
        0.03,
        0.97,
        metric_text(metrics),
        transform=ax.transAxes,
        va="top",
        ha="left",
        bbox={"boxstyle": "round", "facecolor": "white", "alpha": 0.85},
    )
    fig.tight_layout()
    fig.savefig(png_path, dpi=300, bbox_inches="tight")
    plt.close(fig)

    csv_path = png_path.with_suffix(".csv")
    with csv_path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(["actual", "predicted"])
        writer.writerows(zip(targets.tolist(), predictions.tolist()))

    print(f"Saved last-epoch regression plot to: {png_path}")
    print(f"Saved last-epoch predictions to: {csv_path}")


def main():
    args = parse_args()
    if len(set(args.seeds)) != len(args.seeds):
        raise ValueError(f"Duplicate seeds are not allowed: {args.seeds}")
    if args.precision == "fp32":
        torch.set_float32_matmul_precision("highest")
        if torch.cuda.is_available():
            torch.backends.cuda.matmul.allow_tf32 = False
            torch.backends.cudnn.allow_tf32 = False

    save_dir = Path(args.save_dir)
    save_dir.mkdir(parents=True, exist_ok=True)
    device = torch.device(args.device)
    common = trainer_args(args.device, args.precision)
    vision = frozen_vision(args.model, args.checkpoint)

    summary, histories, visual_dim = [], defaultdict(list), None
    for task, (train_arg, test_arg) in TASK_PATH_ARGS.items():
        train_json = getattr(args, train_arg)
        test_json = getattr(args, test_arg)
        print(f"\nTask: {task}")

        train_data = get_ibtracs_reg_dataset(
            train_json, args.ibtracs_npz_dir, args.batch_size, args.workers
        )
        test_data = get_ibtracs_reg_dataset(
            test_json, args.ibtracs_npz_dir, args.batch_size, args.workers
        )
        train_image_loader = unwrap_loader(train_data)
        test_image_loader = unwrap_loader(test_data)

        train_x, train_y = extract_image_features(
            vision,
            train_image_loader,
            device,
            args.precision,
            description=f"Extracting {task} train features",
        )
        test_x, test_y = extract_image_features(
            vision,
            test_image_loader,
            device,
            args.precision,
            description=f"Extracting {task} test features",
        )
        validate_xy(train_x, train_y, f"{task}/train")
        validate_xy(test_x, test_y, f"{task}/test")
        if train_x.shape[1] != test_x.shape[1]:
            raise ValueError(f"{task}: train/test embedding dimensions differ.")

        visual_dim = visual_dim or int(train_x.shape[1])
        if train_x.shape[1] != visual_dim:
            raise ValueError(f"{task}: embedding dim differs from the previous task.")
        print(f"Visual embedding dim: {visual_dim}")

        train_x = l2_normalize(train_x, f"{task}/train")
        test_x = l2_normalize(test_x, f"{task}/test")
        print(
            "Applied F.normalize(features, p=2, dim=1); "
            "each image feature row now has L2 norm 1."
        )

        vision = vision.cpu()
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

        mean, std = target_stats(train_y, task)
        print(f"Training-target mean={mean:.6f}, std={std:.6f} (unbiased=False)")
        train_ds = RegressionFeatureDataset(train_x, train_y)
        test_ds = RegressionFeatureDataset(test_x, test_y)
        task_dir = safe_name(task)

        for seed in args.seeds:
            pl.seed_everything(seed, workers=True)
            train_loader = feature_loader(
                train_ds, args.probe_batch_size, True, seed, args.device == "cuda"
            )
            test_loader = feature_loader(
                test_ds, args.probe_batch_size, False, seed, args.device == "cuda"
            )
            probe = LinearRegressionProbe(
                visual_dim, mean, std, args.lr, args.weight_decay, task, seed
            )
            trainer = pl.Trainer(**common, max_epochs=args.epochs)
            trainer.fit(
                probe,
                train_dataloaders=train_loader,
                val_dataloaders=test_loader,
            )
            if len(probe.history) != args.epochs:
                raise RuntimeError(
                    f"Expected {args.epochs} epoch records, got {len(probe.history)}."
                )
            if probe.last_predictions is None or probe.last_targets is None:
                raise RuntimeError("Last-epoch predictions were not retained.")

            final = probe.history[-1]
            summary.append(
                {
                    "task": task,
                    "seed": seed,
                    "mae": final["test_mae"],
                    "rmse": final["test_rmse"],
                    "r2": final["test_r2"],
                    "pearson": final["test_pearson"],
                }
            )
            histories[seed].extend({"task": task, **row} for row in probe.history)

            run_name = f"seed_{seed}_epoch_{final['epoch']}"
            head_path = save_dir / "linear_heads" / task_dir / f"{run_name}.pt"
            plot_path = save_dir / "regression_plots" / task_dir / f"{run_name}.png"
            plot_metrics = {
                "mae": final["test_mae"],
                "rmse": final["test_rmse"],
                "r2": final["test_r2"],
                "pearson": final["test_pearson"],
            }
            save_linear_head(head_path, probe, visual_dim, final)
            save_regression_plot(
                probe.last_predictions,
                probe.last_targets,
                plot_metrics,
                plot_path,
                task,
                title=f"{task} | seed={seed} | epoch={final['epoch']}",
            )

            print(
                f"FINAL {task} | seed={seed} | last_epoch={final['epoch']} "
                f"mae={final['test_mae']:.4f} rmse={final['test_rmse']:.4f} "
                f"r2={final['test_r2']:.4f} pearson={final['test_pearson']:.4f}"
            )
            del trainer, probe, train_loader, test_loader
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()

        del train_data, test_data, train_image_loader, test_image_loader
        del train_ds, test_ds, train_x, train_y, test_x, test_y
        gc.collect()

    del vision
    write_excel(save_dir / "regression_probe_results.xlsx", summary, dict(histories))


if __name__ == "__main__":
    main()
