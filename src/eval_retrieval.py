import os
import argparse
import torch
import json
from tqdm.auto import tqdm

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from meteolip import create_model_and_transforms, get_tokenizer, get_input_dtype
from meteolip_train.data import get_retrieval_val_dataset
from meteolip_train.precision import get_autocast


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--model", type=str, required=True)
    parser.add_argument("--checkpoint", type=str, required=True)
    parser.add_argument("--val-data", type=str, required=True)
    parser.add_argument("--val-dir-npz", type=str, required=True)
    parser.add_argument("--val-dir-cap", type=str, required=True)
    parser.add_argument("--batch-size", type=int, default=64)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--device", type=str, default="cuda")
    parser.add_argument("--precision", type=str, default="amp")
    parser.add_argument("--topk", type=int, default=5)
    parser.add_argument("--caption-style", type=str, default="natural_event_focused")
    parser.add_argument("--output-dir", type=str, default="ret_results")
    parser.add_argument("--record-source", type=str, nargs="+", default=None)
    return parser.parse_args()


def load_checkpoint(model, checkpoint_path, device):
    ckpt = torch.load(checkpoint_path, map_location=device)

    if isinstance(ckpt, dict):
        if "state_dict" in ckpt:
            state_dict = ckpt["state_dict"]
        elif "model" in ckpt:
            state_dict = ckpt["model"]
        else:
            state_dict = ckpt
    else:
        state_dict = ckpt

    cleaned_state_dict = {}
    for k, v in state_dict.items():
        if k.startswith("module."):
            cleaned_state_dict[k[len("module."):]] = v
        else:
            cleaned_state_dict[k] = v

    cleaned_state_dict = {key.replace(".vat", ".mettok"): val for key, val in cleaned_state_dict.items()}

    missing, unexpected = model.load_state_dict(cleaned_state_dict, strict=True)
    print(f"Loaded checkpoint: {checkpoint_path}")
    print(f"Missing keys: {len(missing)}")
    print(f"Unexpected keys: {len(unexpected)}")


def collect_features_and_metadata(model, dataloader, args):
    device = torch.device(args.device)
    autocast = get_autocast(args.precision, device_type=device.type)
    input_dtype = get_input_dtype(args.precision)

    all_image_features = []
    all_text_features = []
    all_folder_names = []
    all_captions = []

    with torch.inference_mode():
        for batch in tqdm(dataloader, desc="Extracting features"):
            images, texts, folder_names, raw_texts = batch

            images = images.to(device=device, dtype=input_dtype, non_blocking=True)
            texts = texts.to(device=device, non_blocking=True)

            with autocast():
                model_out = model(images, texts)

                if isinstance(model_out, dict):
                    image_features = model_out["image_features"]
                    text_features = model_out["text_features"]
                else:
                    image_features = model_out[0]
                    text_features = model_out[1]

            all_image_features.append(image_features.cpu())
            all_text_features.append(text_features.cpu())

            all_folder_names.extend([str(x) for x in folder_names])
            all_captions.extend([str(x) for x in raw_texts])

    all_image_features = torch.cat(all_image_features, dim=0)
    all_text_features = torch.cat(all_text_features, dim=0)

    return all_image_features, all_text_features, all_folder_names, all_captions


def compute_metrics_from_ranks(ranks):
    ranks = torch.tensor(ranks, dtype=torch.float32)
    metrics = {
        "R@1": (ranks <= 1).float().mean().item(),
        "R@5": (ranks <= 5).float().mean().item(),
        "R@10": (ranks <= 10).float().mean().item(),
        "median_rank": ranks.median().item(),
        "mean_rank": ranks.mean().item(),
    }
    return metrics

@torch.inference_mode()
def run_retrieval(
    image_features,
    text_features,
    captions,
    folder_names,
    logit_scale,
    topk=10,
    chunk_size=512,
):
    if len(image_features) != len(text_features):
        raise ValueError("This function assumes one text per image.")

    num_samples = image_features.shape[0]
    device = image_features.device
    k = min(topk, num_samples)

    # Retrieval metrics should preferably be calculated in FP32
    image_features = image_features.float()
    text_features = text_features.float()

    image_features = torch.nn.functional.normalize(
        image_features,
        dim=-1,
    )
    text_features = torch.nn.functional.normalize(
        text_features,
        dim=-1,
    )

    if torch.is_tensor(logit_scale):
        scale = logit_scale.detach().float().to(device)
    else:
        scale = torch.tensor(
            float(logit_scale),
            device=device,
            dtype=torch.float32,
        )

    true_scores = scale * (image_features * text_features).sum(dim=1)

    image_to_text_ranks = torch.empty(
        num_samples,
        dtype=torch.int64,
        device="cpu",
    )

    text_to_image_better_counts = torch.zeros(
        num_samples,
        dtype=torch.int64,
        device=device,
    )

    topk_indices = torch.empty(
        (num_samples, k),
        dtype=torch.int64,
        device="cpu",
    )

    topk_scores = torch.empty(
        (num_samples, k),
        dtype=torch.float32,
        device="cpu",
    )

    text_features_t = text_features.t().contiguous()

    for start in tqdm(
        range(0, num_samples, chunk_size),
        desc="Computing retrieval",
    ):
        end = min(start + chunk_size, num_samples)
        batch_size = end - start

        scores = scale * (
            image_features[start:end] @ text_features_t
        )

        local_indices = torch.arange(batch_size, device=device)
        global_indices = torch.arange(start, end, device=device)

        # Ensure each diagonal score is identical to true_scores
        scores[local_indices, global_indices] = true_scores[global_indices]

        # Image -> text ranks
        i2t_ranks = (
            (scores > true_scores[start:end, None]).sum(dim=1) + 1
        )

        image_to_text_ranks[start:end] = i2t_ranks.cpu()

        # Text -> image ranks
        text_to_image_better_counts += (
            scores > true_scores[None, :]
        ).sum(dim=0, dtype=torch.int64)

        # Only retrieve the top-k candidates
        indices = torch.topk(
            scores,
            k=k,
            dim=1,
            largest=True,
            sorted=True,
        ).indices

        topk_indices[start:end] = indices.cpu()

        topk_scores[start:end] = torch.gather(
            scores,
            dim=1,
            index=indices,
        ).cpu()

    text_to_image_ranks = (
        text_to_image_better_counts + 1
    ).cpu()

    topk_cosine_similarities = topk_scores / scale.detach().cpu()

    i2t_rank_list = image_to_text_ranks.tolist()
    t2i_rank_list = text_to_image_ranks.tolist()
    topk_index_list = topk_indices.tolist()
    topk_cosine_similarity_list = topk_cosine_similarities.tolist()

    results = []

    for i, indices in enumerate(topk_index_list):
        results.append({
            "folder_name": folder_names[i],
            "rank": i2t_rank_list[i],
            "true_caption": captions[i],
            "top_captions": [
                captions[j] for j in indices
            ],
            "top_folder_names": [
                folder_names[j] for j in indices
            ],
            "top_cosine_similarities": topk_cosine_similarity_list[i],
        })

    metrics = {
        "image_to_text": compute_metrics_from_ranks(
            i2t_rank_list
        ),
        "text_to_image": compute_metrics_from_ranks(
            t2i_rank_list
        ),
    }

    return results, metrics


def print_metrics(metrics):
    headers = ["R@1", "R@5", "R@10", "Median Rank", "Mean Rank"]

    print("\nImage-to-Text metrics:")
    print("\t".join(headers))
    print("\t".join([
        f"{metrics['image_to_text']['R@1']:.4f}",
        f"{metrics['image_to_text']['R@5']:.4f}",
        f"{metrics['image_to_text']['R@10']:.4f}",
        f"{metrics['image_to_text']['median_rank']:.2f}",
        f"{metrics['image_to_text']['mean_rank']:.2f}",
    ]))

    print("\nText-to-Image metrics:")
    print("\t".join(headers))
    print("\t".join([
        f"{metrics['text_to_image']['R@1']:.4f}",
        f"{metrics['text_to_image']['R@5']:.4f}",
        f"{metrics['text_to_image']['R@10']:.4f}",
        f"{metrics['text_to_image']['median_rank']:.2f}",
        f"{metrics['text_to_image']['mean_rank']:.2f}",
    ]))


def save_results_xlsx(results, xlsx_path, num_top_captions=10):
    wb = Workbook()
    ws = wb.active
    ws.title = "retrieval_results"

    ws.append([
        "folder_name",
        "rank",
        "true_caption",
        f"top{num_top_captions}_captions",
        f"top{num_top_captions}_folder_names",
        f"top{num_top_captions}_cosine_similarities",
    ])

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for item in results:
        top_caps = item["top_captions"][:num_top_captions]
        top_folders = item["top_folder_names"][:num_top_captions]
        top_cosine_similarities = item["top_cosine_similarities"][:num_top_captions]

        top_caps_text = "\n".join(
            [f"{i + 1}. {cap}" for i, cap in enumerate(top_caps)]
        )
        top_folders_text = "\n".join(
            [f"{i + 1}. {name}" for i, name in enumerate(top_folders)]
        )

        top_cosine_similarities_text = "\n".join(
            [
                f"{i + 1}. {similarity:.6f}"
                for i, similarity in enumerate(top_cosine_similarities)
            ]
        )

        ws.append([
            item["folder_name"],
            item["rank"],
            item["true_caption"],
            top_caps_text,
            top_folders_text,
            top_cosine_similarities_text,
        ])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 100
    ws.column_dimensions["E"].width = 35
    ws.column_dimensions["F"].width = 30

    ws.freeze_panes = "A2"

    wb.save(xlsx_path)
    print(f"\nSaved XLSX to {xlsx_path}")


def main():
    args = parse_args()

    assert args.caption_style in (
        "natural_event_focused", "impact_focused", "physical_framing", "atmospheric_pattern",
    )

    device = torch.device(args.device)

    model, _, preprocess_val = create_model_and_transforms(
        args.model,
        pretrained=None,
        device=device,
        precision=args.precision,
    )
    load_checkpoint(model, args.checkpoint, device)
    model.eval()

    tokenizer = get_tokenizer(args.model)

    data = {}
    data["val"] = get_retrieval_val_dataset(
        args.val_data, args.val_dir_npz, args.val_dir_cap,
        caption_style=args.caption_style, return_metadata=True,
        batch_size=args.batch_size, workers=args.workers,
        tokenizer=tokenizer,
        record_source=args.record_source,
    )

    print("Number of test samples:", len(data["val"].dataloader.dataset))
    dataset_names = sorted({
        item[0].split("__", 1)[0]
        for item in data["val"].dataloader.dataset.data
    })
    print("Dataset names:", dataset_names)

    if args.record_source is not None:
        assert set(dataset_names) == set(args.record_source), (
            f"Dataset names do not match record_source. "
            f"Found: {dataset_names}, expected: {args.record_source}"
        )

    image_features, text_features, folder_names, captions = collect_features_and_metadata(
        model=model,
        dataloader=data["val"].dataloader,
        args=args,
    )

    if len(image_features) != len(captions):
        raise ValueError(
            f"Feature count and caption count do not match: "
            f"{len(image_features)} vs {len(captions)}"
        )

    if len(folder_names) != len(captions):
        raise ValueError(
            f"folder_names and captions count do not match: "
            f"{len(folder_names)} vs {len(captions)}"
        )

    logit_scale = model.logit_scale.exp().detach().cpu()

    results, metrics = run_retrieval(
        image_features=image_features,
        text_features=text_features,
        captions=captions,
        folder_names=folder_names,
        logit_scale=logit_scale,
        topk=args.topk,
    )

    print_metrics(metrics)

    os.makedirs(args.output_dir, exist_ok=True)
    with open(os.path.join(args.output_dir, f"{args.caption_style}.json"), "w", encoding="utf-8") as f:
        json.dump(metrics, f, ensure_ascii=False, indent=2)
    save_results_xlsx(
        results, os.path.join(args.output_dir, f"{args.caption_style}.xlsx"),
        num_top_captions=args.topk,
    )


if __name__ == "__main__":
    main()